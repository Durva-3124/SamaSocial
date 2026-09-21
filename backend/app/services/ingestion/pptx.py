"""PPTX ingestion using python-pptx."""
import logging
from typing import TYPE_CHECKING

from app.core.errors import AppError
from app.models.chunk import Chunk, Locator
from app.services.chunking import split_text
from app.services.ingestion.base import IngestResult

if TYPE_CHECKING:
    pass

logger = logging.getLogger(__name__)

_MAX_SLIDE_WORDS = 350


def _shape_text(shape) -> str:
    """Recursively extract text from a shape, including group shapes."""
    from pptx.util import Pt  # noqa: F401
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "\n".join(_shape_text(s) for s in shape.shapes if _shape_text(s))

    parts: list[str] = []

    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)

    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        rows: list[str] = []
        for row in shape.table.rows:
            cells = " | ".join(cell.text.strip() for cell in row.cells)
            if cells.strip(" |"):
                rows.append(cells)
        if rows:
            parts.append("\n".join(rows))

    return "\n".join(parts)


def _slide_text(slide, slide_num: int) -> str:
    """Build the full text block for a slide."""
    parts: list[str] = []

    # Title first
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        title = slide.shapes.title.text_frame.text.strip()
        if title:
            parts.append(f"Title: {title}")

    # Body shapes (skip title, already handled)
    for shape in slide.shapes:
        if shape == slide.shapes.title:
            continue
        text = _shape_text(shape)
        if text:
            parts.append(text)

    # Speaker notes
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            parts.append(f"Speaker notes: {notes}")

    return "\n".join(parts)


def ingest_pptx(source_id: str, filename: str, data: bytes) -> IngestResult:
    """Ingest a PPTX file into chunks with slide locators."""
    try:
        import io
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise AppError("INVALID_PPTX", "The file is not a valid PPTX.", 422) from exc

    chunks: list[Chunk] = []
    image_only_slides: list[int] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        text = _slide_text(slide, slide_num)
        if not text.strip():
            image_only_slides.append(slide_num)
            continue

        pieces = split_text(text, max_words=_MAX_SLIDE_WORDS) if len(text.split()) > _MAX_SLIDE_WORDS else [text]
        for piece in pieces:
            chunks.append(
                Chunk(
                    source_id=source_id,
                    source_type="pptx",
                    text=piece,
                    locator=Locator(slide=slide_num),
                )
            )

    warnings: list[str] = []
    if image_only_slides:
        nums = ", ".join(str(n) for n in image_only_slides)
        warnings.append(
            f"{len(image_only_slides)} slides contain only images and were skipped (slides: {nums})."
        )

    if not chunks:
        raise AppError("NO_TEXT", "No text could be extracted from this presentation.", 422)

    return IngestResult(
        name=filename,
        source_type="pptx",
        chunks=chunks,
        warnings=warnings,
    )
