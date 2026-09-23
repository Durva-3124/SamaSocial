"""Fixture builders for the evaluation harness.

All functions are pure (no network, no LLM).  They return bytes / strings
that can be fed directly into the ingestion layer.
"""
from __future__ import annotations

import json
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Inches, Pt

# ── Fact registry ─────────────────────────────────────────────────────────────
FACTS: list[dict] = [
    # PDF — Photosynthesis (pages 1-5)
    {"id": "f01", "source": "pdf", "locator": "page 1",
     "text": "Photosynthesis converts light energy into chemical energy stored in glucose.",
     "keywords": ["photosynthesis", "glucose", "chemical energy"]},
    {"id": "f02", "source": "pdf", "locator": "page 1",
     "text": "The overall equation for photosynthesis is: 6CO2 + 6H2O + light produces C6H12O6 + 6O2.",
     "keywords": ["6CO2", "6H2O", "C6H12O6"]},
    {"id": "f03", "source": "pdf", "locator": "page 2",
     "text": "Chlorophyll is the primary pigment that absorbs light, mainly in the red and blue wavelengths.",
     "keywords": ["chlorophyll", "pigment", "red", "blue"]},
    {"id": "f04", "source": "pdf", "locator": "page 2",
     "text": "The light-dependent reactions occur in the thylakoid membranes of the chloroplast.",
     "keywords": ["thylakoid", "light-dependent", "chloroplast"]},
    {"id": "f05", "source": "pdf", "locator": "page 3",
     "text": "The Calvin cycle (light-independent reactions) takes place in the stroma.",
     "keywords": ["Calvin cycle", "stroma", "light-independent"]},
    {"id": "f06", "source": "pdf", "locator": "page 3",
     "text": "ATP and NADPH produced in the light reactions power the Calvin cycle.",
     "keywords": ["ATP", "NADPH", "Calvin cycle"]},
    {"id": "f07", "source": "pdf", "locator": "page 4",
     "text": "Carbon fixation is the process by which CO2 is incorporated into organic molecules via RuBisCO.",
     "keywords": ["carbon fixation", "RuBisCO", "CO2"]},
    {"id": "f08", "source": "pdf", "locator": "page 4",
     "text": "The rate of photosynthesis increases with light intensity up to a saturation point.",
     "keywords": ["rate", "light intensity", "saturation"]},
    {"id": "f09", "source": "pdf", "locator": "page 5",
     "text": "Photorespiration occurs when RuBisCO fixes O2 instead of CO2, reducing photosynthetic efficiency.",
     "keywords": ["photorespiration", "RuBisCO", "O2"]},
    {"id": "f10", "source": "pdf", "locator": "page 5",
     "text": "C4 plants like maize use a two-stage carbon fixation to minimise photorespiration.",
     "keywords": ["C4", "maize", "photorespiration"]},
    {"id": "f11", "source": "pdf", "locator": "page 5",
     "text": "CAM plants open stomata at night to reduce water loss during carbon fixation.",
     "keywords": ["CAM", "stomata", "night", "water loss"]},
    {"id": "f12", "source": "pdf", "locator": "page 5",
     "text": "Oxygen produced during photosynthesis comes from the splitting of water molecules.",
     "keywords": ["oxygen", "splitting", "water"]},

    # PPTX — REST APIs (slides 1-6)
    {"id": "f13", "source": "pptx", "locator": "slide 1",
     "text": "REST stands for Representational State Transfer, an architectural style for distributed systems.",
     "keywords": ["REST", "Representational State Transfer", "architectural"]},
    {"id": "f14", "source": "pptx", "locator": "slide 1",
     "text": "REST was defined by Roy Fielding in his 2000 doctoral dissertation.",
     "keywords": ["Roy Fielding", "2000", "dissertation"]},
    {"id": "f15", "source": "pptx", "locator": "slide 2",
     "text": "The six REST constraints are: client-server, stateless, cacheable, uniform interface, layered system, and code on demand.",
     "keywords": ["stateless", "cacheable", "uniform interface", "layered"]},
    {"id": "f16", "source": "pptx", "locator": "slide 2",
     "text": "Statelessness means each request must contain all information needed to process it.",
     "keywords": ["stateless", "request", "information"]},
    {"id": "f17", "source": "pptx", "locator": "slide 3",
     "text": "HTTP methods used in REST: GET (read), POST (create), PUT (replace), PATCH (update), DELETE (remove).",
     "keywords": ["GET", "POST", "PUT", "PATCH", "DELETE"]},
    {"id": "f18", "source": "pptx", "locator": "slide 3",
     "text": "A 200 OK response indicates success; 201 Created is returned after a successful POST.",
     "keywords": ["200 OK", "201 Created", "POST"]},
    {"id": "f19", "source": "pptx", "locator": "slide 4",
     "text": "Resources in REST are identified by URIs; representations can be JSON or XML.",
     "keywords": ["URI", "JSON", "XML", "resources"]},
    {"id": "f20", "source": "pptx", "locator": "slide 4",
     "text": "HATEOAS (Hypermedia as the Engine of Application State) allows clients to navigate APIs via links.",
     "keywords": ["HATEOAS", "hypermedia", "links"]},
    {"id": "f21", "source": "pptx", "locator": "slide 5",
     "text": "Rate limiting protects APIs from abuse; the 429 Too Many Requests status code signals this.",
     "keywords": ["rate limiting", "429", "Too Many Requests"]},
    {"id": "f22", "source": "pptx", "locator": "slide 5",
     "text": "API versioning strategies include URI versioning (/v1/), header versioning, and query parameter versioning.",
     "keywords": ["versioning", "/v1/", "header"]},
    {"id": "f23", "source": "pptx", "locator": "slide 6",
     "text": "OpenAPI Specification (formerly Swagger) is the standard for documenting REST APIs.",
     "keywords": ["OpenAPI", "Swagger", "documenting"]},
    {"id": "f24", "source": "pptx", "locator": "slide 6",
     "text": "JWT (JSON Web Token) is commonly used for stateless authentication in REST APIs.",
     "keywords": ["JWT", "JSON Web Token", "authentication"]},

    # HTML — Machine Learning basics
    # locator matches Chunk.locator_text() output: f'section "{heading}"'
    {"id": "f25", "source": "web", "locator": 'section "Introduction"',
     "text": "Machine learning is a subset of artificial intelligence that enables systems to learn from data.",
     "keywords": ["machine learning", "artificial intelligence", "data"]},
    {"id": "f26", "source": "web", "locator": 'section "Supervised Learning"',
     "text": "Supervised learning uses labelled training data to learn a mapping from inputs to outputs.",
     "keywords": ["supervised", "labelled", "inputs", "outputs"]},
    {"id": "f27", "source": "web", "locator": 'section "Unsupervised Learning"',
     "text": "Unsupervised learning finds hidden patterns in unlabelled data, such as clustering and dimensionality reduction.",
     "keywords": ["unsupervised", "unlabelled", "clustering"]},
    {"id": "f28", "source": "web", "locator": 'section "Overfitting"',
     "text": "Overfitting occurs when a model learns noise in the training data and performs poorly on new data.",
     "keywords": ["overfitting", "noise", "training data"]},
    {"id": "f29", "source": "web", "locator": 'section "Regularisation"',
     "text": "Regularisation techniques like L1 (Lasso) and L2 (Ridge) penalise large weights to reduce overfitting.",
     "keywords": ["regularisation", "L1", "Lasso", "L2", "Ridge"]},
    {"id": "f30", "source": "web", "locator": 'section "Cross-Validation"',
     "text": "K-fold cross-validation splits data into K subsets and trains K models, each tested on a different fold.",
     "keywords": ["K-fold", "cross-validation", "subsets"]},
    {"id": "f31", "source": "web", "locator": 'section "Neural Networks"',
     "text": "A neural network consists of layers of interconnected nodes (neurons) that transform inputs via activation functions.",
     "keywords": ["neural network", "neurons", "activation functions"]},
    {"id": "f32", "source": "web", "locator": 'section "Gradient Descent"',
     "text": "Gradient descent minimises the loss function by iteratively adjusting weights in the direction of the negative gradient.",
     "keywords": ["gradient descent", "loss function", "weights"]},

    # YouTube transcript — Python programming
    {"id": "f33", "source": "youtube", "locator": "at 0:00",
     "text": "Python is an interpreted, high-level, general-purpose programming language created by Guido van Rossum.",
     "keywords": ["Python", "interpreted", "Guido van Rossum"]},
    {"id": "f34", "source": "youtube", "locator": "at 1:00",
     "text": "Python uses indentation to define code blocks instead of curly braces.",
     "keywords": ["indentation", "code blocks", "curly braces"]},
    {"id": "f35", "source": "youtube", "locator": "at 2:00",
     "text": "Python lists are mutable ordered sequences; tuples are immutable ordered sequences.",
     "keywords": ["lists", "mutable", "tuples", "immutable"]},
    {"id": "f36", "source": "youtube", "locator": "at 3:00",
     "text": "Dictionaries in Python store key-value pairs and are implemented as hash tables.",
     "keywords": ["dictionaries", "key-value", "hash tables"]},
    {"id": "f37", "source": "youtube", "locator": "at 4:00",
     "text": "Python's GIL (Global Interpreter Lock) prevents true multi-threading for CPU-bound tasks.",
     "keywords": ["GIL", "Global Interpreter Lock", "multi-threading"]},
    {"id": "f38", "source": "youtube", "locator": "at 5:00",
     "text": "List comprehensions provide a concise way to create lists: x times 2 for x in range 10.",
     "keywords": ["list comprehensions", "concise", "range"]},
    {"id": "f39", "source": "youtube", "locator": "at 6:00",
     "text": "Python decorators are functions that modify the behaviour of other functions using the @ syntax.",
     "keywords": ["decorators", "functions", "@"]},
    {"id": "f40", "source": "youtube", "locator": "at 7:00",
     "text": "The __init__ method is the constructor in Python classes, called when an object is instantiated.",
     "keywords": ["__init__", "constructor", "instantiated"]},
]


# ── PDF builder ───────────────────────────────────────────────────────────────

def build_photosynthesis_pdf() -> bytes:
    """Build a 5-page PDF about photosynthesis containing facts f01-f12."""
    pages: dict[int, list[dict]] = {
        1: [FACTS[0], FACTS[1]],
        2: [FACTS[2], FACTS[3]],
        3: [FACTS[4], FACTS[5]],
        4: [FACTS[6], FACTS[7]],
        5: [FACTS[8], FACTS[9], FACTS[10], FACTS[11]],
    }
    doc = fitz.open()
    for page_num in range(1, 6):
        page = doc.new_page()
        y = 72.0
        page.insert_text((50, y), f"Photosynthesis — Page {page_num}", fontsize=14)
        y += 30
        for fact in pages[page_num]:
            page.insert_text((50, y), fact["text"], fontsize=11)
            y += 50
    data = doc.tobytes()
    doc.close()
    return data


# ── PPTX builder ──────────────────────────────────────────────────────────────

def build_rest_api_pptx() -> bytes:
    """Build a 6-slide PPTX about REST APIs containing facts f13-f24."""
    slides_content: dict[int, tuple[str, list[dict]]] = {
        1: ("REST APIs — Introduction", [FACTS[12], FACTS[13]]),
        2: ("REST Constraints", [FACTS[14], FACTS[15]]),
        3: ("HTTP Methods and Status Codes", [FACTS[16], FACTS[17]]),
        4: ("Resources and HATEOAS", [FACTS[18], FACTS[19]]),
        5: ("Rate Limiting and Versioning", [FACTS[20], FACTS[21]]),
        6: ("Documentation and Auth", [FACTS[22], FACTS[23]]),
    }
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    for slide_num in range(1, 7):
        title_text, facts = slides_content[slide_num]
        slide = prs.slides.add_slide(blank_layout)
        tb_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
        tb_title.text_frame.text = title_text
        tb_title.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        for i, fact in enumerate(facts):
            tb = slide.shapes.add_textbox(
                Inches(0.5), Inches(1.2 + i * 1.4), Inches(9), Inches(1.2)
            )
            tb.text_frame.word_wrap = True
            tb.text_frame.text = fact["text"]
            tb.text_frame.paragraphs[0].runs[0].font.size = Pt(13)

    import io
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ── HTML builder ──────────────────────────────────────────────────────────────

def build_ml_html() -> str:
    """Build an HTML page about ML basics containing facts f25-f32."""
    sections = [
        ("Introduction", [FACTS[24]]),
        ("Supervised Learning", [FACTS[25]]),
        ("Unsupervised Learning", [FACTS[26]]),
        ("Overfitting", [FACTS[27]]),
        ("Regularisation", [FACTS[28]]),
        ("Cross-Validation", [FACTS[29]]),
        ("Neural Networks", [FACTS[30]]),
        ("Gradient Descent", [FACTS[31]]),
    ]
    body = ""
    for heading, facts in sections:
        body += f"<h2>{heading}</h2>\n"
        for fact in facts:
            body += f"<p>{fact['text']}</p>\n"
    return (
        "<!DOCTYPE html>\n<html><head><title>Machine Learning Basics</title></head>\n"
        f"<body>\n<h1>Machine Learning Basics</h1>\n{body}</body></html>"
    )


# ── YouTube transcript builder ────────────────────────────────────────────────

def build_python_transcript() -> list[dict]:
    """Build a fake transcript [{text, start, duration}] about Python (facts f33-f40)."""
    return [
        {"text": fact["text"], "start": float(i * 60), "duration": 55.0}
        for i, fact in enumerate(FACTS[32:40])
    ]


# ── facts.json writer ─────────────────────────────────────────────────────────

def write_facts(path: Path) -> None:
    """Write FACTS list to a JSON file."""
    path.write_text(json.dumps(FACTS, indent=2))
