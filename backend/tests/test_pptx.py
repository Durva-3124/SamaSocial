"""Tests for PPTX ingestion — presentations built in-test with python-pptx."""
import io

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.core.errors import AppError
from app.services.ingestion.pptx import ingest_pptx


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _pptx_bytes(build_fn) -> bytes:
    prs = Presentation()
    build_fn(prs)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_text_slide(prs: Presentation, title: str, body: str) -> None:
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = body


def _add_blank_slide(prs: Presentation) -> None:
    layout = prs.slide_layouts[6]  # Blank
    prs.slides.add_slide(layout)


# ---------------------------------------------------------------------------
# Happy path
# ---------------------------------------------------------------------------

def test_title_and_body_extracted() -> None:
    data = _pptx_bytes(lambda prs: _add_text_slide(prs, "My Title", "Body content here."))
    result = ingest_pptx("s1", "deck.pptx", data)
    assert result.source_type == "pptx"
    assert result.name == "deck.pptx"
    combined = " ".join(c.text for c in result.chunks)
    assert "My Title" in combined
    assert "Body content here" in combined


def test_slide_numbers_correct() -> None:
    def build(prs):
        _add_text_slide(prs, "Slide One", "Content A")
        _add_text_slide(prs, "Slide Two", "Content B")
        _add_text_slide(prs, "Slide Three", "Content C")

    data = _pptx_bytes(build)
    result = ingest_pptx("s", "f.pptx", data)
    slide_nums = {c.locator.slide for c in result.chunks}
    assert slide_nums == {1, 2, 3}


def test_table_content_extracted() -> None:
    def build(prs):
        layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        rows, cols = 2, 3
        table = slide.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(6), Inches(2)).table
        table.cell(0, 0).text = "Header1"
        table.cell(0, 1).text = "Header2"
        table.cell(0, 2).text = "Header3"
        table.cell(1, 0).text = "Val1"
        table.cell(1, 1).text = "Val2"
        table.cell(1, 2).text = "Val3"

    data = _pptx_bytes(build)
    result = ingest_pptx("s", "f.pptx", data)
    combined = " ".join(c.text for c in result.chunks)
    assert "Header1" in combined
    assert "Val2" in combined


def test_speaker_notes_extracted() -> None:
    def build(prs):
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Title"
        slide.placeholders[1].text = "Body"
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = "These are the speaker notes."

    data = _pptx_bytes(build)
    result = ingest_pptx("s", "f.pptx", data)
    combined = " ".join(c.text for c in result.chunks)
    assert "speaker notes" in combined.lower()
    assert "These are the speaker notes" in combined


def test_image_only_slide_skipped_with_warning() -> None:
    def build(prs):
        _add_text_slide(prs, "Real Slide", "Some content here.")
        _add_blank_slide(prs)  # image-only (blank)

    data = _pptx_bytes(build)
    result = ingest_pptx("s", "f.pptx", data)
    slide_nums = {c.locator.slide for c in result.chunks}
    assert 2 not in slide_nums
    assert len(result.warnings) == 1
    assert "skipped" in result.warnings[0].lower()


def test_all_chunks_have_correct_source_id() -> None:
    data = _pptx_bytes(lambda prs: _add_text_slide(prs, "T", "B"))
    result = ingest_pptx("my-id", "f.pptx", data)
    assert all(c.source_id == "my-id" for c in result.chunks)


# ---------------------------------------------------------------------------
# Error cases
# ---------------------------------------------------------------------------

def test_all_image_slides_raises_no_text() -> None:
    data = _pptx_bytes(lambda prs: (_add_blank_slide(prs), _add_blank_slide(prs)))
    with pytest.raises(AppError) as exc_info:
        ingest_pptx("s", "f.pptx", data)
    assert exc_info.value.code == "NO_TEXT"


def test_corrupt_bytes_raises_invalid_pptx() -> None:
    with pytest.raises(AppError) as exc_info:
        ingest_pptx("s", "f.pptx", b"not a pptx file")
    assert exc_info.value.code == "INVALID_PPTX"
